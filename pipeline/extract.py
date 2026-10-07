"""Steg 2: Trekk ut tekst per side, gjenkjenn kapittel/post, og bygg
kapitteloppslag fra tallgrunnlaget til Gul bok."""
import json
import re

import pymupdf

from common import CACHE, DATA, PAGES_DIR, PDF_DIR, read_json, write_json

# «kap. 1320 post 30», «kap. 1320, post 30», «kapittel 1320», «kap. 1320 postane 30 og 31»
REF = re.compile(
    r"\b[Kk]ap(?:ittel|\.)\s?(\d{1,4})\b(?![–-]\d)"
    r"(?:\s*,?\s*(?:under\s+)?post(?:ane|ene)?\.?\s*(\d{2})\b(?:\s*(?:og|,|–|-)\s*(\d{2})\b)?)?"
)
# Kapitteloverskrift i Prop. 1 S: en linje som starter med «Kap. 1320 Statens vegvesen»
HEAD = re.compile(r"^Kap\.\s?(\d{1,4})\s+([A-ZÆØÅ].{2,80})$")


def load_chapters():
    """{kap: {"navn", "dep", "poster": {post: {"navn", "belop"}}}} fra tallgrunnlaget."""
    path = CACHE / "tallgrunnlag.xlsx"
    if not path.exists():
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[-1]
    rows = ws.iter_rows(values_only=True)
    head = [str(h).strip() for h in next(rows)]
    ix = {h: i for i, h in enumerate(head)}
    chapters = {}
    for r in rows:
        if r[ix["kap_nr"]] is None:
            continue
        kap = str(int(r[ix["kap_nr"]]))
        post = f"{int(r[ix['post_nr']]):02d}"
        c = chapters.setdefault(kap, {"navn": str(r[ix["kap_navn"]]).strip(),
                                      "dep": str(r[ix["fdep_navn"]]).strip(), "poster": {}})
        c["poster"][post] = {"navn": str(r[ix["post_navn"]]).strip(), "belop": r[ix["beløp"]]}
    return chapters


HYPH = re.compile(r"(\w)[-­]\n(?!(?:og|eller|til|samt)\b)([a-zæøå])")


def page_lines(page):
    """[(tekst, størrelse, fet)] per linje."""
    out = []
    for b in page.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            spans = [s for s in l["spans"] if s["text"].strip()]
            if not spans:
                continue
            t = re.sub(r"[ \t ]+", " ", "".join(s["text"] for s in l["spans"])).strip()
            s0 = spans[0]
            bold = "bold" in s0["font"].lower() or bool(s0["flags"] & 16)
            out.append((t, round(s0["size"], 1), bold))
    return out


def running_headers(all_lines):
    """Linjer som går igjen øverst/nederst på mange sider (topptekst/bunntekst)."""
    from collections import Counter
    c = Counter()
    for lines in all_lines:
        edge = {t for t, _, _ in lines[:5] + lines[-3:]}
        c.update(edge)
    n = max(len(all_lines), 1)
    return {t for t, k in c.items() if k >= max(3, 0.25 * n) and not re.fullmatch(r"\d+", t)}


def extract_doc(doc, chapters):
    path = PDF_DIR / doc["cache_file"]
    with pymupdf.open(path) as pdf:
        all_lines = [page_lines(p) for p in pdf]
    headers = running_headers(all_lines)
    # Overskriftsstørrelse for kapitler: størrelsen på fete «Kap. NNNN»-linjer
    head_sizes = [sz for lines in all_lines for t, sz, b in lines if b and HEAD.match(t)]
    head_size = max(set(head_sizes), key=head_sizes.count) if head_sizes else None

    out, context = [], None
    for i, lines in enumerate(all_lines, start=1):
        label = None
        for t, _, _ in lines[:4] + lines[-3:]:
            if re.fullmatch(r"\d{1,4}", t):
                label = int(t)
                break
        body_lines = [t for t, _, _ in lines
                      if t not in headers and not (label is not None and t == str(label))]
        # Kapittelet fra forrige side gjelder bare videre hvis denne siden har en
        # «Post NN»-overskrift eller nevner kapittelnummeret; ellers er vi usikre.
        if context and not any(b and re.match(r"^Post \d{2}\b", t) for t, _, b in lines) \
                and not re.search(rf"\b{context}\b", " ".join(body_lines)):
            context = None
        heads, page_context = [], context
        for t, sz, b in lines:
            if head_size and b and sz >= head_size - 0.2:
                m = HEAD.match(t)
                if m and m.group(1) in chapters:
                    heads.append(m.group(1))
                    context = m.group(1)
                elif len(t) > 3 and not re.fullmatch(r"[\d\s.,–-]+", t):
                    context = None  # ny seksjon på samme nivå som ikke er et kapittel
        if heads:
            page_context = heads[0]
        text = HYPH.sub(r"\1\2", "\n".join(body_lines))
        refs = []
        for m in REF.finditer(text):
            kap = m.group(1)
            if kap not in chapters:
                continue
            for post in (m.group(2), m.group(3)):
                if post and post in chapters[kap]["poster"]:
                    refs.append(f"{kap}.{post}")
            refs.append(kap)
        refs = list(dict.fromkeys(refs))
        out.append({"doc": doc["id"], "page": i, "label": label, "text": text,
                    "kap_context": page_context if head_size else None,
                    "kap_refs": refs[:12]})
    with open(PAGES_DIR / f"{doc['id']}.jsonl", "w", encoding="utf-8") as f:
        for p in out:
            f.write(json.dumps(p, ensure_ascii=False) + "\n")
    return out


def extract(only=None):
    manifest = read_json(DATA / "manifest.json")
    chapters = load_chapters()
    if chapters:
        write_json(DATA / "kapitler.json", chapters)
    total = 0
    for doc in manifest["documents"]:
        if only and doc["type"] not in only:
            continue
        pages = extract_doc(doc, chapters)
        empty = sum(1 for p in pages if len(p["text"]) < 40)
        total += len(pages)
        print(f"  {doc['id']}: {len(pages)} sider ({empty} nesten tomme)")
    print(f"Ferdig: {total} sider, {len(chapters)} kapitler i oppslaget.")


if __name__ == "__main__":
    extract()
